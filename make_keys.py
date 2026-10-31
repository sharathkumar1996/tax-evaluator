"""Answer-key workbooks: blank template, worked examples, and drafting a key from a correct ATX return."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from grader_core import KEY_COLS, F1040_FIELDS, ReturnPDF, describe, guess_category

F = "Arial"
NAVY = "1F3864"
HDR = PatternFill("solid", fgColor=NAVY)
INPUT = PatternFill("solid", fgColor="FFF2CC")
BLUE = Font(name=F, color="0000FF")
BODY = Font(name=F)
BOLD = Font(name=F, bold=True)
WIDTHS = [6, 22, 13, 18, 40, 18, 8, 11, 70]

HOW = [
    ("How to fill in an answer key", None),
    ("", None),
    ("Quickest way", "In the app, use 'Build key from a correct return': upload the trainer's correct ATX return and you get "
                     "this workbook with every filled-in item already listed. Then set Marks, Categories and Feedback."),
    ("One file = one case", "Fill in Case ID, Case Title, Tax Year and 'Marks deducted per unexpected entry' at the top."),
    ("Edit", "Only the yellow cells (blue text). Do not rename the sheet or the column headers."),
    ("Step", "1, 2, 3 ... order shown in the student report."),
    ("Category", "Group for category-wise marks, e.g. Personal Information, Filing Status, Dependents, Income, "
                 "Schedule C, Deductions, Tax & Credits, Payments & Refund."),
    ("Form", "Form or schedule as printed: 1040, Schedule 1, Schedule C, Schedule SE, Form 8863 ... "
             "For a second copy of the same schedule write 'Schedule C #2'."),
    ("Line - numbered lines", "Line number exactly as printed: 1, 2b, 25a, 31. Lettered lines (e.g. Schedule C line B) also work."),
    ("Line - Form 1040 fields", "FS (filing status), first_name, last_name, ssn, spouse_first_name, spouse_last_name, spouse_ssn, "
                                "address, apt, city, state, zip, foreign_country, virtual_currency (Yes/No), occupation, "
                                "spouse_occupation, you_as_dependent, spouse_as_dependent, you_born_before, you_blind, "
                                "spouse_born_before, spouse_blind (Checked/Unchecked)."),
    ("Line - dependents", "dep1_name, dep1_ssn, dep1_relationship, dep1_ctc, dep1_odc (Checked/Unchecked); same for dep2 ... dep4."),
    ("Line - any other box", "text: <label printed on the form>   e.g.  text: Principal business or profession\n"
                             "check: <label printed next to the box>   e.g.  check: Cash   (add  | near: <words on the same row>  "
                             "if the label appears more than once)."),
    ("Correct Value", "Amounts as whole numbers (48000, not $48,000; negative -500). Text exactly as it should appear. "
                      "Filing status S, MFJ, MFS, HOH or QW. Checkboxes: Checked / Unchecked. Leave blank if the box must be empty."),
    ("Marks", "Marks for this step. Any difference from the key loses the full marks for the step."),
    ("Tolerance ($)", "Allowed difference for rounding only. Usually 1 for amounts, 0 for everything else."),
    ("Feedback if Wrong", "Shown in the student's report when the step is wrong. Explain the correct treatment."),
    ("Unexpected entries", "If 'Marks deducted per unexpected entry' has a number, every amount the student entered on a line "
                           "that is NOT in this key (a line that should be blank) is listed as a mistake and deducts that many "
                           "marks. Leave it blank to switch the check off; 0 lists them without deducting."),
]

CATS = ("Personal Information,Filing Status,Dependents,Income,Adjustments,Deductions,Tax & Credits,"
        "Payments & Refund,Schedule 1,Schedule 2,Schedule 3,Schedule A,Schedule B,Schedule C,Schedule D,Schedule E,"
        "Schedule SE")


def build(case_id, title, year, rows, unexpected=1, blank_rows=0, path=None):
    """rows: list of (step, category, form, line, description, correct, marks, tolerance, feedback). Returns bytes."""
    wb = Workbook()
    k = wb.active
    k.title = "Answer Key"
    k["A1"] = "ANSWER KEY"
    k["A1"].font = Font(name=F, bold=True, size=14, color=NAVY)
    info = [("Case ID", case_id), ("Case Title", title), ("Tax Year", year),
            ("Marks deducted per unexpected entry", unexpected)]
    for r, (lab, val) in enumerate(info, 2):
        k.cell(r, 1, lab).font = BOLD
        c = k.cell(r, 2, val)
        c.font, c.fill = BLUE, INPUT
    k.cell(5, 3, "<- blank = don't check lines outside the key").font = Font(name=F, italic=True, color="595959")
    HR = 7
    for i, (c, w) in enumerate(zip(KEY_COLS, WIDTHS), 1):
        cell = k.cell(HR, i, c)
        cell.font = Font(name=F, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        k.column_dimensions[cell.column_letter].width = w
    k.column_dimensions["A"].width = 34   # also fits the info labels
    k.row_dimensions[HR].height = 30
    k.freeze_panes = k.cell(HR + 1, 1)
    r = HR + 1
    for row in rows:
        for c, v in enumerate(row, 1):
            cell = k.cell(r, c, v)
            cell.font, cell.fill = BLUE, INPUT
            cell.alignment = Alignment(wrap_text=c in (5, 9), vertical="top")
            if c == 6 and isinstance(v, (int, float)):
                cell.number_format = "#,##0;-#,##0;0"
        r += 1
    for _ in range(blank_rows):
        for c in range(1, len(KEY_COLS) + 1):
            k.cell(r, c).fill, k.cell(r, c).font = INPUT, BLUE
        r += 1
    first, last = HR + 1, r - 1
    k.cell(r + 1, 5, "Total marks").font = BOLD
    k.cell(r + 1, 7, f"=SUM(G{first}:G{last})").font = BOLD
    dv = DataValidation(type="list", allow_blank=True, formula1=f'"{CATS}"')
    dv.error, dv.errorStyle, dv.showErrorMessage = "Pick from the list or type your own category.", "warning", True
    k.add_data_validation(dv)
    dv.add(f"B{first}:B{max(last, first) + 100}")

    h = wb.create_sheet("How to Use")
    h.column_dimensions["A"].width = 26
    h.column_dimensions["B"].width = 115
    for i, (a, b) in enumerate(HOW, 1):
        h.cell(i, 1, a).font = BOLD
        h.cell(i, 1).alignment = Alignment(vertical="top")
        if b:
            c = h.cell(i, 2, b)
            c.font = BODY
            c.alignment = Alignment(wrap_text=True, vertical="top")
    h["A1"].font = Font(name=F, bold=True, size=14, color=NAVY)
    buf = io.BytesIO()
    wb.save(buf)
    if path:
        open(path, "wb").write(buf.getvalue())
    return buf.getvalue()


def _sort_key(fld):
    order = ["FS"] + list(F1040_FIELDS)
    ln = str(fld["line"])
    form_rank = 0 if fld["form"] == "1040" else 1
    if ln in order:
        return (form_rank, fld["form"], 0, order.index(ln), "")
    if ln.startswith("dep"):
        return (form_rank, fld["form"], 1, 0, ln)
    import re
    m = re.match(r"^(\d+)([a-z]*)$", ln.lower())
    return (form_rank, fld["form"], 2, int(m.group(1)) if m else 999, m.group(2) if m else ln)


def draft_from_return(pdf_bytes, case_id, title, year, marks=1, unexpected=1):
    """Read every filled-in item from a CORRECT return and lay it out as an answer key."""
    ret = ReturnPDF(pdf_bytes)
    rows = []
    for i, fld in enumerate(sorted(ret.all_fields(), key=_sort_key), 1):
        v = fld["value"]
        if isinstance(v, float) and v == int(v):
            v = int(v)
        num = isinstance(v, (int, float))
        rows.append((i, guess_category(fld["form"], fld["line"]), fld["form"], fld["line"],
                     describe(fld["form"], fld["line"], f"{fld['form']} line {fld['line']}"),
                     v, marks, 1 if num else 0, ""))
    return build(case_id, title, year, rows, unexpected), len(rows), ret.forms_found()


if __name__ == "__main__":
    build("C00", "Case title here", 2020,
          [(1, "Income", "1040", "1", "Wages, salaries, tips", 48000, 5, 1,
            "EXAMPLE ROW - replace. Line 1 is W-2 box 1 wages.")], unexpected=1, blank_rows=59,
          path="answer_key_template.xlsx")
    print("template built")
