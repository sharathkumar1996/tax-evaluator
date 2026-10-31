"""
grader_core.py - reads an ATX return PDF (Form 1040 + any schedules/forms) and
grades it against an Excel answer key.

Reading a return
----------------
* Each page is assigned to a form from its header text:
  "SCHEDULE C (Form 1040)" -> SCH C,  "Form 8863" -> 8863,  1040 pages -> 1040.
  A page whose header says "Page 2" etc. continues the previous copy of that form;
  otherwise it starts a new copy (so a second Schedule C is "SCH C #2").
* On a form's pages, a line is found by its bold line number (e.g. "31", "2b", "25a").
  The student's figure is the filled-in amount immediately to its right on the
  same row. ATX prints filled-in data at ~8.8pt; the form's own text is smaller,
  which separates student entries from printed form text.
* A blank box reads as 0. A form that is not in the PDF reads as "missing".
* Named Form 1040 fields (names, SSNs, address, checkboxes, dependents) and any
  "text: <printed label>" / "check: <printed label>" field are read from the box
  under/next to the printed label.
* Grading is strict: a figure either matches the key (within the rounding tolerance) or
  loses the step's marks. Optional: entries on lines the key says nothing about
  ("unexpected entries") can each deduct marks.
"""
import io
import re
from dataclasses import dataclass, field

import pdfplumber
from openpyxl import load_workbook

# ------------------------------------------------------------------ settings
DATA_MIN_SIZE = 8.3      # pt: ATX filled-in data is ~8.78pt, form text <= ~8pt
ROW_TOL = 6              # pt: amount may sit a few points above/below its label
HEADER_FRACTION = 0.16   # top part of the page used to identify the form

MONEY = re.compile(r"^\(?-?\$?\d{1,3}(,\d{3})*\)?$|^\(?-?\$?\d+\)?$")
LABEL = re.compile(r"^\d{1,3}[a-z]{0,2}$")

FILING_STATUS = {
    "single": "S", "s": "S",
    "married filing jointly": "MFJ", "mfj": "MFJ",
    "married filing separately": "MFS", "mfs": "MFS",
    "head of household": "HOH", "hoh": "HOH",
    "qualifying widow(er)": "QW", "qualifying widow": "QW", "qw": "QW", "qss": "QW",
    "qualifying surviving spouse": "QW",
}


# ------------------------------------------------------------ form names
def norm_form(name):
    """'Schedule C' / 'Sch C' / 'SCH C #2' / 'Form 8863' / '1040' -> canonical id."""
    s = str(name or "").strip().upper()
    inst = 1
    m = re.search(r"#\s*(\d+)\s*$", s)
    if m:
        inst = int(m.group(1))
        s = s[: m.start()].strip()
    s = re.sub(r"^FORM\s+", "", s)
    m = re.match(r"^(SCHEDULE|SCHED|SCH)\.?\s*([A-Z0-9-]+)$", s)
    if m:
        s = "SCH " + m.group(2)
    if s in ("1040-SR", "1040SR"):
        s = "1040"
    return s, inst


def norm_line(line):
    s = str(line or "").strip().lower()
    s = re.sub(r"^line\s*", "", s)
    if s.endswith(".0"):          # Excel turned 31 into 31.0
        s = s[:-2]
    return s


def _identify(page):
    """Return (form_id, is_continuation) for a page, or (None, False)."""
    h = page.height * HEADER_FRACTION
    try:
        head = page.crop((0, 0, page.width, h)).extract_text() or ""
    except Exception:
        head = ""
    full = (page.extract_text() or "")[:400]
    for txt in (head, full):
        cont = bool(re.search(r"\bPage\s+\d\b", txt))
        m = re.search(r"\bSCHEDULE\s+([A-Z0-9]{1,4}(?:-[A-Z]{1,3})?)\b[\s\S]{0,140}?\(Form\s+1040", txt, re.I)
        if m:
            return "SCH " + m.group(1).upper(), cont
        if "U.S. Individual Income Tax Return" in txt or re.search(r"\bForm\s+1040(-SR)?\s*\(\d{4}\)", txt):
            return "1040", cont
        m = re.search(r"(?:^|\n)\s*Form\s+(W-2G?|\d{3,4}[A-Z]{0,2}(?:-[A-Z]{1,3})?)\b", txt)
        if m and not m.group(1).startswith("1040"):
            return m.group(1).upper(), cont
    return None, False


# ------------------------------------------------------------- tokenising
def _tokens(page):
    """Split a page into tokens: bold line numbers vs regular text/amounts.
    Spaces and dot leaders are dropped; a change of boldness starts a new token."""
    chars = [c for c in page.chars if c["text"].strip() and c["text"] != "\u00a0"
             and not (c["text"] == "." and c["size"] < DATA_MIN_SIZE)]      # drop dot leaders, keep decimals
    chars.sort(key=lambda c: (round(c["top"]), c["x0"]))
    toks = []
    for c in chars:
        bold = "Bold" in c["fontname"] or "Black" in c["fontname"]
        t = toks[-1] if toks else None
        if (t and abs(t["top"] - c["top"]) <= 1.5 and -0.5 <= c["x0"] - t["x1"] <= 2.5
                and t["bold"] == bold):
            t["text"] += c["text"]
            t["x1"] = c["x1"]
            t["size"] = max(t["size"], c["size"])
        else:
            toks.append({"text": c["text"], "x0": c["x0"], "x1": c["x1"], "top": c["top"],
                         "bottom": c["bottom"], "bold": bold, "size": c["size"]})
    for t in toks:
        t["cy"] = (t["top"] + t["bottom"]) / 2
    return toks


def _to_number(tok):
    t = tok.replace("$", "").replace(",", "")
    neg = t.startswith("(") or t.startswith("-")
    t = t.strip("()-")
    try:
        v = float(t)
    except ValueError:
        return None
    return -v if neg else v


def _filing_status(page):
    words = sorted(page.extract_words(x_tolerance=1.5), key=lambda w: (w["top"], w["x0"]))
    rows = []
    for w in words:
        if rows and abs(rows[-1][0]["top"] - w["top"]) <= 3:
            rows[-1].append(w)
        else:
            rows.append([w])
    for ws in rows:
        ws.sort(key=lambda w: w["x0"])
        text = " ".join(w["text"] for w in ws)
        if "Single" not in text and "Filing Status" not in text:
            continue
        for x in [w for w in ws if w["text"] in ("X", "x", "✓", "✔")]:
            nxt = " ".join(w["text"] for w in ws if w["x0"] > x["x1"])[:40].lower()
            for label, code in FILING_STATUS.items():
                if len(label) > 3 and nxt.startswith(label):
                    return code
    return None


# ---------------------------------------------------------------- reader
def _norm_txt(t):
    return re.sub(r"[^a-z0-9]", "", str(t).lower())


class _PageText:
    """Printed form text (small font) as rows of characters, plus filled-in data words (ATX ~8.8pt)."""

    def __init__(self, page):
        self.width = page.width
        chars = [c for c in page.chars if c["size"] < DATA_MIN_SIZE and c["text"].strip()]
        chars.sort(key=lambda c: (round(c["top"]), c["x0"]))
        self.rows = []
        for c in chars:
            if self.rows and abs(self.rows[-1][0]["top"] - c["top"]) <= 2.5:
                self.rows[-1].append(c)
            else:
                self.rows.append([c])
        for r in self.rows:
            r.sort(key=lambda c: c["x0"])
        self.norm = ["".join(_norm_txt(c["text"]) for c in r) for r in self.rows]
        self.idx = [[i for i, c in enumerate(r) for _ in _norm_txt(c["text"])] for r in self.rows]
        words = page.extract_words(x_tolerance=1.5, keep_blank_chars=True, extra_attrs=["size", "fontname"])
        self.data = [w for w in words if w["size"] >= DATA_MIN_SIZE and "Bold" not in w["fontname"]
                     and w["text"].strip()]
        for w in self.data:
            w["text"] = w["text"].strip()
            w["cy"] = (w["top"] + w["bottom"]) / 2

    def find(self, phrase, near=None):
        """All boxes (x0, x1, top, bottom) where the printed phrase occurs; optional 'near' text on the same row."""
        p = _norm_txt(phrase)
        out = []
        if not p:
            return out
        for ri, rn in enumerate(self.norm):
            if near:
                ctx = "".join(self.norm[j] for j in range(len(self.rows))
                              if abs(self.rows[j][0]["top"] - self.rows[ri][0]["top"]) <= 6)
                if _norm_txt(near) not in ctx:
                    continue
            start = rn.find(p)
            while start >= 0:
                r = self.rows[ri]
                a, b = r[self.idx[ri][start]], r[self.idx[ri][start + len(p) - 1]]
                out.append({"x0": a["x0"], "x1": b["x1"], "top": min(a["top"], b["top"]),
                            "bottom": max(a["bottom"], b["bottom"]), "row": ri})
                start = rn.find(p, start + 1)
        return out

    def _next_label_x(self, box):
        """x where the next printed label starts to the right on the same row (box boundary)."""
        r = self.rows[box["row"]]
        gaps = [c["x0"] for i, c in enumerate(r) if c["x0"] > box["x1"] + 6
                and (i == 0 or c["x0"] - r[i - 1]["x1"] > 6)]
        return min(gaps, default=self.width)

    def text_under(self, phrase, near=None):
        boxes = self.find(phrase, near)
        if not boxes:
            return None
        b = boxes[0]
        right = self._next_label_x(b)
        got = [w for w in self.data if b["bottom"] - 1 <= w["top"] <= b["bottom"] + 16
               and b["x0"] - 4 <= w["x0"] < right - 4]
        if not got:   # value printed on the same row, to the right of the label
            got = [w for w in self.data if abs(w["cy"] - (b["top"] + b["bottom"]) / 2) <= 5
                   and b["x1"] < w["x0"] < right]
        got.sort(key=lambda w: w["x0"])
        return " ".join(w["text"] for w in got)

    def checked(self, phrase, near=None):
        boxes = self.find(phrase, near)
        if not boxes:
            return None
        for b in boxes:
            cy = (b["top"] + b["bottom"]) / 2
            for w in self.data:
                if w["text"] in ("X", "x", "✓", "✔") and abs(w["cy"] - cy) <= 5 \
                        and b["x0"] - 18 <= w["x1"] <= b["x0"] + 1:
                    return True
        return False


# Named fields on Form 1040 that are not numbered lines -> how to read them
F1040_FIELDS = {
    "first_name": ("text", 1, "Your first name and middle initial", None),
    "last_name": ("text", 1, "Last name", "Your first name"),
    "ssn": ("text", 1, "Your social security number", None),
    "spouse_first_name": ("text", 1, "If joint return, spouse's first name", None),
    "spouse_last_name": ("text", 1, "Last name", "spouse's first name"),
    "spouse_ssn": ("text", 1, "Spouse's social security number", None),
    "address": ("text", 1, "Home address (number and street)", None),
    "apt": ("text", 1, "Apt. no.", None),
    "city": ("text", 1, "City, town, or post office", None),
    "state": ("text", 1, "State", "ZIP code"),
    "zip": ("text", 1, "ZIP code", None),
    "foreign_country": ("text", 1, "Foreign country name", None),
    "virtual_currency": ("yesno", 1, "virtual currency", None),
    "occupation": ("text", 2, "Your occupation", None),
    "spouse_occupation": ("text", 2, "Spouse's occupation", None),
    "you_as_dependent": ("check", 1, "You as a dependent", None),
    "spouse_as_dependent": ("check", 1, "Your spouse as a dependent", None),
    "you_born_before": ("check", 1, "Were born before January 2", None),
    "you_blind": ("check", 1, "Are blind", None),
    "spouse_born_before": ("check", 1, "Was born before January 2", None),
    "spouse_blind": ("check", 1, "Is blind", None),
}
DEP_FIELDS = ("name", "ssn", "relationship", "ctc", "odc")
FIELD_CATEGORY = {"fs": "Filing Status", "dep": "Dependents"}


@dataclass
class FormCopy:
    form: str
    instance: int
    pages: list = field(default_factory=list)      # page numbers (1-based)
    tokens: list = field(default_factory=list)     # per page token lists
    texts: list = field(default_factory=list)      # per page _PageText


class ReturnPDF:
    def __init__(self, data: bytes):
        self.copies = {}           # (form, instance) -> FormCopy
        self.page_map = []         # [(page_no, form label)]
        self.filing_status = None
        self._deps = None
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            self.n_pages = len(pdf.pages)
            counts, last = {}, None
            for i, page in enumerate(pdf.pages, 1):
                form, cont = _identify(page)
                if form is None:
                    self.page_map.append((i, "(not a graded form)"))
                    last = None
                    continue
                if cont and last and last.form == form:
                    copy = last
                else:
                    counts[form] = counts.get(form, 0) + 1
                    copy = FormCopy(form, counts[form])
                    self.copies[(form, copy.instance)] = copy
                copy.pages.append(i)
                copy.tokens.append(_tokens(page))
                copy.texts.append(_PageText(page))
                last = copy
                lab = form if copy.instance == 1 else f"{form} #{copy.instance}"
                self.page_map.append((i, lab))
                if form == "1040" and self.filing_status is None:
                    self.filing_status = _filing_status(page)

    def forms_found(self):
        return sorted({f if i == 1 else f"{f} #{i}" for (f, i) in self.copies})

    # ---- dependents table on Form 1040 page 1
    def dependents(self):
        if self._deps is not None:
            return self._deps
        self._deps = []
        c = self.copies.get(("1040", 1))
        if not c:
            return self._deps
        pt = c.texts[0]
        try:
            name = pt.find("(1) First name")[0]
            ssn = pt.find("(2) Social security")[0]
            rel = pt.find("(3) Relationship")[0]
            ctc = pt.find("Child tax credit")[0]
            odc = pt.find("Credit for other dependents")[0]
        except IndexError:
            return self._deps
        line1 = [t for t in c.tokens[0] if t["bold"] and t["text"] == "1" and t["x0"] > 0.6 * pt.width]
        y0 = max(name["bottom"], ctc["bottom"], odc["bottom"]) + 1
        y1 = line1[0]["top"] - 4 if line1 else y0 + 60
        words = sorted([w for w in pt.data if y0 <= w["top"] <= y1], key=lambda w: (w["top"], w["x0"]))
        rows = []
        for w in words:
            if rows and abs(rows[-1][0]["top"] - w["top"]) <= 4:
                rows[-1].append(w)
            else:
                rows.append([w])
        for r in rows:
            pick = lambda a, b: " ".join(w["text"] for w in sorted(r, key=lambda w: w["x0"]) if a <= w["x0"] < b)
            is_x = lambda a, b: any(w["text"] in ("X", "x", "✓", "✔") and a <= w["x0"] < b for w in r)
            self._deps.append({
                "name": pick(name["x0"] - 6, ssn["x0"] - 6), "ssn": pick(ssn["x0"] - 6, rel["x0"] - 6),
                "relationship": pick(rel["x0"] - 6, min(ctc["x0"], odc["x0"]) - 14),
                "ctc": "Checked" if is_x(ctc["x0"] - 16, odc["x0"] - 16) else "Unchecked",
                "odc": "Checked" if is_x(odc["x0"] - 16, pt.width) else "Unchecked"})
        return self._deps

    # ---- numbered line lookup
    @staticmethod
    def _col_x(toks, lab, col):
        """x-centre of a column heading such as '(h)' or 'A' printed above the row."""
        want = col.lower().strip("()")
        heads = [t for t in toks if t["text"].lower().strip("()") == want and t["text"].strip("()") != t["text"] or
                 (t["text"].lower() == want and len(want) == 1 and t["size"] < DATA_MIN_SIZE)]
        heads = [t for t in heads if 0 < lab["top"] - t["top"] < 260]
        if not heads:
            return None
        h = min(heads, key=lambda t: lab["top"] - t["top"])
        return (h["x0"] + h["x1"]) / 2

    def _line(self, copy, ln):
        """(value, printed) for a numbered/lettered line on a form copy; (None, False) if no such line.
        ln may name a column: '1a (h)' or '3 (A)'."""
        col = None
        m = re.match(r"^(.+?)\s*\(([a-z])\)$", ln)
        if m:
            ln, col = m.group(1).strip(), m.group(2)
        best = None
        for toks in copy.tokens:
            for lab in [t for t in toks if t["bold"] and t["text"].lower() == ln]:
                row_labels = [t for t in toks if t["bold"] and LABEL.match(t["text"].lower())
                              and abs(t["cy"] - lab["cy"]) <= ROW_TOL and t["x0"] > lab["x1"] + 1]
                stop = min([t["x0"] for t in row_labels], default=1e9)
                data = [t for t in toks if not t["bold"] and t["size"] >= DATA_MIN_SIZE
                        and lab["x1"] < t["x0"] < stop and abs(t["cy"] - lab["cy"]) <= ROW_TOL]
                amts = [t for t in data if MONEY.match(t["text"])]
                if col:
                    cx = self._col_x(toks, lab, col)
                    if cx is None:
                        continue
                    amts = [t for t in amts if abs((t["x0"] + t["x1"]) / 2 - cx) < 45 or t["x0"] - 6 <= cx <= t["x1"] + 6]
                    data = amts
                if amts:
                    pick = min(amts, key=lambda t: t["x0"]) if not col else \
                        min(amts, key=lambda t: abs((t["x0"] + t["x1"]) / 2 - cx))
                    score, v, printed = 3 + lab["x0"] / 1000, _to_number(pick["text"]), True
                elif data:
                    data.sort(key=lambda t: t["x0"])
                    score, v, printed = 2 + lab["x0"] / 1000, " ".join(t["text"] for t in data), True
                else:
                    score, v, printed = 1 + lab["x0"] / 1000, 0.0, False
                if best is None or score > best[0]:
                    best = (score, v, printed)
        return (None, False) if best is None else (best[1], best[2])

    def value(self, form, line):
        """Return (value, note). value None = could not read / form missing."""
        f, inst = norm_form(form)
        raw = str(line or "").strip()
        ln = norm_line(raw)
        copy = self.copies.get((f, inst))
        name = f if inst == 1 else f"{f} #{inst}"
        if copy is None:
            return None, f"{name} not found in the return"
        if f == "1040" and ln in ("fs", "filing status", "filing_status"):
            return (self.filing_status, "") if self.filing_status else (None, "Filing status box not read")
        m = re.match(r"^dep(\d)_(\w+)$", ln)
        if f == "1040" and m and m.group(2) in DEP_FIELDS:
            deps = self.dependents()
            i = int(m.group(1)) - 1
            if i >= len(deps):
                return ("Unchecked" if m.group(2) in ("ctc", "odc") else ""), ""
            return deps[i][m.group(2)], ""
        if f == "1040" and ln in F1040_FIELDS:
            kind, pg, phrase, near = F1040_FIELDS[ln]
            if pg > len(copy.texts):
                return None, f"Page {pg} of Form 1040 not found"
            pt = copy.texts[pg - 1]
            if kind == "text":
                v = pt.text_under(phrase, near)
            elif kind == "check":
                c = pt.checked(phrase, near)
                v = None if c is None else ("Checked" if c else "Unchecked")
            else:
                y, n = pt.checked("Yes", phrase), pt.checked("No", phrase)
                v = None if y is None and n is None else ("Yes" if y else "No" if n else "")
            return (v, "") if v is not None else (None, f"'{phrase}' not found on Form 1040")
        m = re.match(r"^(text|check):\s*(.+?)(?:\s*\|\s*near:\s*(.+))?$", raw, re.I)
        if m:
            kind, phrase, near = m.group(1).lower(), m.group(2), m.group(3)
            for pt in copy.texts:
                if kind == "text":
                    v = pt.text_under(phrase, near)
                else:
                    c = pt.checked(phrase, near)
                    v = None if c is None else ("Checked" if c else "Unchecked")
                if v is not None:
                    return v, ""
            return None, f"'{phrase}' not found on {name}"
        v, _ = self._line(copy, ln)
        if v is None:
            return None, f"Line {ln} not found on {name}"
        return v, ""

    # ---- everything filled in (used to draft keys and to catch unexpected entries)
    def all_fields(self):
        out = []
        for (f, inst), copy in self.copies.items():
            fname = f if inst == 1 else f"{f} #{inst}"
            if f == "1040" and inst == 1:
                if self.filing_status:
                    out.append({"form": fname, "line": "FS", "desc": "Filing status", "value": self.filing_status})
                for key in F1040_FIELDS:
                    v, _ = self.value("1040", key)
                    if v not in (None, "", "Unchecked"):
                        out.append({"form": fname, "line": key, "desc": key.replace("_", " ").capitalize(), "value": v})
                for i, d in enumerate(self.dependents(), 1):
                    for k_ in DEP_FIELDS:
                        out.append({"form": fname, "line": f"dep{i}_{k_}", "desc": f"Dependent {i} {k_}",
                                    "value": d[k_]})
            seen = set()
            for pi, toks in enumerate(copy.tokens):
                pt = copy.texts[pi]
                for lab in toks:
                    ln = lab["text"].lower()
                    if not lab["bold"] or not LABEL.match(ln) or ln in seen:
                        continue
                    seen.add(ln)
                    cols = self._row_columns(toks, lab)
                    if cols:
                        for cname in cols:
                            v, printed = self._line(copy, f"{ln} ({cname})")
                            if printed:
                                out.append({"form": fname, "line": f"{lab['text']} ({cname})", "value": v, "desc": ""})
                        continue
                    v, printed = self._line(copy, ln)
                    if not printed:
                        continue
                    out.append({"form": fname, "line": lab["text"], "value": v,
                                "desc": self._describe(pt, lab)})
        return out

    def _row_columns(self, toks, lab):
        """Column headings for a row that carries 2+ amounts (e.g. Schedule D 1a: d, e, g, h)."""
        row_labels = [t for t in toks if t["bold"] and LABEL.match(t["text"].lower())
                      and abs(t["cy"] - lab["cy"]) <= ROW_TOL and t["x0"] > lab["x1"] + 1]
        stop = min([t["x0"] for t in row_labels], default=1e9)
        amts = [t for t in toks if not t["bold"] and t["size"] >= DATA_MIN_SIZE and MONEY.match(t["text"])
                and lab["x1"] < t["x0"] < stop and abs(t["cy"] - lab["cy"]) <= ROW_TOL]
        if len(amts) < 2:
            return []
        heads = [t for t in toks if re.fullmatch(r"\([a-z]\)", t["text"]) and 0 < lab["top"] - t["top"] < 260]
        names = []
        for a in amts:
            ac = (a["x0"] + a["x1"]) / 2
            near = [h for h in heads if abs((h["x0"] + h["x1"]) / 2 - ac) < 45]
            if near:
                h = min(near, key=lambda h: (lab["top"] - h["top"]) + abs((h["x0"] + h["x1"]) / 2 - ac))
                names.append(h["text"].strip("()"))
        return list(dict.fromkeys(names))

    @staticmethod
    def _describe(pt, lab):
        cy = (lab["top"] + lab["bottom"]) / 2
        best = []
        for r in pt.rows:
            if abs((r[0]["top"] + r[0]["bottom"]) / 2 - cy) <= 4:
                best += [c for c in r if c["x1"] < lab["x0"] - 2]
        best.sort(key=lambda c: c["x0"])
        txt, prev = "", None
        for c in best:
            if prev is not None and c["x0"] - prev > 1.2:
                txt += " "
            txt += c["text"]
            prev = c["x1"]
        txt = re.sub(r"(\s*\.\s*){2,}", " ", txt)
        txt = re.sub(r"^\s*\d{1,2}[a-z]?\s+", "", txt).strip(" .")
        return txt[:70]


# ------------------------------------------------------------ answer key
KEY_COLS = ["Step", "Category", "Form", "Line", "Description", "Correct Value", "Marks",
            "Tolerance ($)", "Feedback if Wrong"]

LINES_1040_2020 = {
    "1": "Wages, salaries, tips", "2a": "Tax-exempt interest", "2b": "Taxable interest",
    "3a": "Qualified dividends", "3b": "Ordinary dividends", "4a": "IRA distributions", "4b": "IRA distributions - taxable",
    "5a": "Pensions and annuities", "5b": "Pensions and annuities - taxable", "6a": "Social security benefits",
    "6b": "Social security benefits - taxable", "7": "Capital gain or (loss)", "8": "Other income from Schedule 1",
    "9": "Total income", "10a": "Adjustments from Schedule 1", "10b": "Charitable contributions (standard deduction filers)",
    "10c": "Total adjustments to income", "11": "Adjusted gross income", "12": "Standard or itemized deduction",
    "13": "Qualified business income deduction", "14": "Add lines 12 and 13", "15": "Taxable income", "16": "Tax",
    "17": "Amount from Schedule 2, line 3", "18": "Add lines 16 and 17", "19": "Child tax credit / credit for other dependents",
    "20": "Amount from Schedule 3, line 7", "21": "Add lines 19 and 20", "22": "Tax after credits",
    "23": "Other taxes (Schedule 2, line 10)", "24": "Total tax", "25a": "Federal tax withheld - Form(s) W-2",
    "25b": "Federal tax withheld - Form(s) 1099", "25c": "Federal tax withheld - other forms",
    "25d": "Total federal tax withheld", "26": "Estimated tax payments", "27": "Earned income credit",
    "28": "Additional child tax credit", "29": "American opportunity credit", "30": "Recovery rebate credit",
    "31": "Amount from Schedule 3, line 13", "32": "Total other payments and refundable credits",
    "33": "Total payments", "34": "Amount overpaid", "35a": "Refund", "36": "Applied to 2021 estimated tax",
    "37": "Amount you owe", "38": "Estimated tax penalty",
}


def guess_category(form, line):
    f, _ = norm_form(form)
    ln = norm_line(line)
    if f != "1040":
        return ("Schedule " + f[4:]) if f.startswith("SCH ") else f"Form {f}"
    if ln in ("fs", "filing status"):
        return "Filing Status"
    if ln.startswith("dep"):
        return "Dependents"
    if ln in F1040_FIELDS:
        return "Personal Information"
    m = re.match(r"^(\d+)", ln)
    n = int(m.group(1)) if m else 0
    return ("Income" if n <= 9 else "Adjustments" if n <= 11 else "Deductions" if n <= 15
            else "Tax & Credits" if n <= 24 else "Payments & Refund")


def describe(form, line, fallback=""):
    f, _ = norm_form(form)
    ln = norm_line(line)
    if f == "1040":
        if ln in LINES_1040_2020:
            return LINES_1040_2020[ln]
        if ln == "fs":
            return "Filing status"
        m = re.match(r"^dep(\d)_(\w+)$", ln)
        if m:
            return f"Dependent {m.group(1)} - " + {"name": "name", "ssn": "SSN", "relationship": "relationship",
                                                   "ctc": "child tax credit box", "odc": "credit for other dependents box"}[m.group(2)]
        if ln in F1040_FIELDS:
            return FIELD_DESC.get(ln, ln.replace("_", " ").capitalize())
    return fallback


FIELD_DESC = {
    "first_name": "First name and middle initial", "last_name": "Last name", "ssn": "Social security number",
    "spouse_first_name": "Spouse's first name and initial", "spouse_last_name": "Spouse's last name",
    "spouse_ssn": "Spouse's social security number", "address": "Home address", "apt": "Apartment number",
    "city": "City / town", "state": "State", "zip": "ZIP code", "foreign_country": "Foreign country",
    "virtual_currency": "Virtual currency question (Yes/No)", "occupation": "Occupation",
    "spouse_occupation": "Spouse's occupation", "you_as_dependent": "Someone can claim you as a dependent",
    "spouse_as_dependent": "Someone can claim your spouse as a dependent",
    "you_born_before": "You were born before 2 Jan 1956", "you_blind": "You are blind",
    "spouse_born_before": "Spouse was born before 2 Jan 1956", "spouse_blind": "Spouse is blind",
}


class KeyError_(Exception):
    pass


@dataclass
class AnswerKey:
    case_id: str
    title: str
    tax_year: int
    steps: list
    unexpected_marks: float = None    # None = don't check; 0 = list only; >0 = deduct per entry

    @property
    def total(self):
        return sum(s["Marks"] for s in self.steps)

    @property
    def categories(self):
        out = []
        for s in self.steps:
            if s["Category"] not in out:
                out.append(s["Category"])
        return out


def load_key(data: bytes) -> AnswerKey:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    if "Answer Key" not in wb.sheetnames:
        raise KeyError_("The workbook has no sheet named 'Answer Key'. Use the template.")
    ws = wb["Answer Key"]
    info, hdr_row = {}, None
    for r in range(1, 30):
        a = ws.cell(r, 1).value
        if a == "Step" and ws.cell(r, 3).value == "Form":
            hdr_row = r
            break
        if isinstance(a, str):
            info[a.strip().rstrip(":").lower()] = ws.cell(r, 2).value
    if hdr_row is None:
        raise KeyError_("Could not find the table header row (Step, Category, Form, Line ...).")
    hdr = [str(ws.cell(hdr_row, c).value or "").strip() for c in range(1, len(KEY_COLS) + 1)]
    if hdr != KEY_COLS:
        raise KeyError_(f"Column headers don't match the current template. Expected: {', '.join(KEY_COLS)}. "
                        "Download the latest template from the app.")
    steps, problems = [], []
    for r in range(hdr_row + 1, ws.max_row + 1):
        row = dict(zip(KEY_COLS, [ws.cell(r, c).value for c in range(1, len(KEY_COLS) + 1)]))
        if row["Form"] in (None, "") and row["Line"] in (None, ""):
            continue
        where = f"Row {r}"
        if row["Form"] in (None, "") or row["Line"] in (None, ""):
            problems.append(f"{where}: Form and Line are both required.")
            continue
        if row["Correct Value"] is None:
            row["Correct Value"] = ""          # blank = the box must be empty
        try:
            row["Marks"] = float(row["Marks"])
        except (TypeError, ValueError):
            problems.append(f"{where}: Marks must be a number.")
            continue
        row["Tolerance ($)"] = float(row["Tolerance ($)"] or 0)
        row["Category"] = str(row["Category"] or guess_category(row["Form"], row["Line"])).strip()
        row["Description"] = row["Description"] or describe(row["Form"], row["Line"], f"{row['Form']} line {row['Line']}")
        row["Step"] = row["Step"] if row["Step"] is not None else len(steps) + 1
        f, inst = norm_form(row["Form"])
        row["_key"] = (f, inst, norm_line(row["Line"]))
        steps.append(row)
    if problems:
        raise KeyError_("\n".join(problems))
    if not steps:
        raise KeyError_("The answer key has no steps.")
    try:
        year = int(info.get("tax year") or 2020)
    except ValueError:
        year = 2020
    um = info.get("marks deducted per unexpected entry")
    try:
        um = None if um in (None, "") else float(um)
    except ValueError:
        um = None
    return AnswerKey(str(info.get("case id") or "CASE"), str(info.get("case title") or ""), year, steps, um)


# ------------------------------------------------------------- grading
def _is_num(x):
    if isinstance(x, bool):
        return False
    try:
        float(str(x).replace(",", "").replace("$", ""))
        return True
    except (TypeError, ValueError):
        return False


def _num(x):
    return float(str(x).replace(",", "").replace("$", ""))


def _clean_text(x):
    t = re.sub(r"\s+", " ", str(x)).strip().strip(",.;").upper()
    return t


def _match(student, correct, tol):
    if student is None:
        return False
    c, s = correct, student
    if isinstance(c, str) and c.strip() == "":
        return (str(s).strip() == "") or (_is_num(s) and _num(s) == 0) or str(s) == "Unchecked"
    if _is_num(c) and _is_num(s) and not re.search(r"\d-\d", str(c)):
        return abs(_num(s) - _num(c)) <= float(tol or 0) + 1e-9
    cs, ss_ = str(c).strip().lower(), str(s).strip().lower()
    if cs in FILING_STATUS or ss_ in FILING_STATUS:
        return FILING_STATUS.get(cs, cs.upper()) == FILING_STATUS.get(ss_, ss_.upper())
    if re.fullmatch(r"[\d\s-]{9,11}", str(c).strip()):          # SSN / EIN: compare digits
        return re.sub(r"\D", "", str(c)) == re.sub(r"\D", "", str(s))
    if cs in ("x", "checked", "yes-checked"):
        return ss_ in ("x", "checked", "\u2713", "\u2714")
    if cs in ("unchecked", "blank"):
        return ss_ in ("unchecked", "")
    return _clean_text(c) == _clean_text(s)


def grade(ret: ReturnPDF, key: AnswerKey):
    rows = []
    for s in key.steps:
        val, note = ret.value(s["Form"], s["Line"])
        ok = _match(val, s["Correct Value"], s["Tolerance ($)"])
        rows.append({
            "Step": s["Step"], "Category": s["Category"], "Form": s["Form"], "Line": s["Line"],
            "Description": s["Description"], "Correct Value": s["Correct Value"], "Student Value": val,
            "Marks": s["Marks"], "Awarded": s["Marks"] if ok else 0.0,
            "Result": "Correct" if ok else "Wrong",
            "Feedback": "" if ok else (s["Feedback if Wrong"] or ""),
            "Note": note, "Needs Review": val is None and "not found in the return" not in note,
        })
    if key.unexpected_marks is not None:
        keyed = {s["_key"] for s in key.steps}
        n = len(rows)
        for fld in ret.all_fields():
            f, inst = norm_form(fld["form"])
            k_ = (f, inst, norm_line(fld["line"]))
            ln_ = norm_line(fld["line"])
            if k_ in keyed or ln_.startswith("dep") or ln_ in F1040_FIELDS or ln_ == "fs":
                continue            # only amounts on numbered lines count as unexpected entries
            v = fld["value"]
            if v in ("", 0, 0.0, None, "Unchecked"):
                continue
            n += 1
            rows.append({
                "Step": n, "Category": "Unexpected entries", "Form": fld["form"], "Line": fld["line"],
                "Description": describe(fld["form"], fld["line"], f"{fld['form']} line {fld['line']}"),
                "Correct Value": "", "Student Value": v, "Marks": 0.0, "Awarded": -float(key.unexpected_marks),
                "Result": "Unexpected entry",
                "Feedback": "This line should be blank for this case. Remove the entry and check where it came from.",
                "Note": "", "Needs Review": False})
    return rows


def summarise(rows):
    cats = {}
    for r in rows:
        c = cats.setdefault(r["Category"], {"Category": r["Category"], "Awarded": 0.0, "Marks": 0.0,
                                             "Steps": 0, "Mistakes": 0})
        c["Awarded"] += r["Awarded"]
        c["Marks"] += r["Marks"]
        c["Steps"] += 1
        c["Mistakes"] += r["Result"] != "Correct"
    for c in cats.values():
        c["Pct"] = max(0, c["Awarded"]) / c["Marks"] if c["Marks"] else 0
    out_of = sum(r["Marks"] for r in rows)
    total = max(0.0, sum(r["Awarded"] for r in rows))
    return list(cats.values()), total, out_of


FS_NAMES = {"S": "Single", "MFJ": "Married filing jointly", "MFS": "Married filing separately",
            "HOH": "Head of household", "QW": "Qualifying widow(er)"}


def fmt(v):
    if v is None:
        return "Not found"
    if isinstance(v, str) and v.strip().upper() in FS_NAMES:
        return FS_NAMES[v.strip().upper()]
    if _is_num(v) and not isinstance(v, str):
        v = float(v)
        return f"{v:,.0f}" if v == int(v) else f"{v:,.2f}"
    return str(v)
